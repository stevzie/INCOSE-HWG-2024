#!/usr/bin/env python3
"""
Create a sample Excel file for testing the Files API.
This creates a workbook with sample questionnaire data.
"""

import openpyxl
from openpyxl import Workbook

# Create a new workbook
wb = Workbook()

# First sheet - Questionnaire Data
ws1 = wb.active
ws1.title = "Questionnaire"

# Add headers
ws1['A1'] = "Question ID"
ws1['B1'] = "Question Text"
ws1['C1'] = "Category"
ws1['D1'] = "Type"

# Add sample data
questions = [
    ["Q1", "What is your primary role?", "Demographics", "Multiple Choice"],
    ["Q2", "How many years of experience do you have?", "Demographics", "Number"],
    ["Q3", "Required Questions: Please rate your satisfaction", "Feedback", "Scale"],
    ["Q4", "What improvements would you suggest?", "Feedback", "Open Text"],
    ["Q5", "Would you recommend this to others?", "Feedback", "Yes/No"],
]

for idx, question in enumerate(questions, start=2):
    ws1[f'A{idx}'] = question[0]
    ws1[f'B{idx}'] = question[1]
    ws1[f'C{idx}'] = question[2]
    ws1[f'D{idx}'] = question[3]

# Second sheet - Metadata
ws2 = wb.create_sheet("Metadata")
ws2['A1'] = "Property"
ws2['B1'] = "Value"
ws2['A2'] = "Survey Name"
ws2['B2'] = "Sample Survey 2024"
ws2['A3'] = "Version"
ws2['B3'] = "1.0"
ws2['A4'] = "Total Questions"
ws2['B4'] = len(questions)

# Third sheet - Analysis
ws3 = wb.create_sheet("Analysis")
ws3['A1'] = "Summary Statistics"
ws3['A3'] = "Total Responses"
ws3['B3'] = 150
ws3['A4'] = "Completion Rate"
ws3['B4'] = "92%"

# Save the workbook
output_path = "/home/user/INCOSE-HWG-2024/sample_survey.xlsx"
wb.save(output_path)
print(f"Sample Excel file created at: {output_path}")
print(f"  - Sheet 1: {ws1.title} ({len(questions)} questions)")
print(f"  - Sheet 2: {ws2.title}")
print(f"  - Sheet 3: {ws3.title}")
